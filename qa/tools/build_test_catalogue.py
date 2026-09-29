#!/usr/bin/env python3
"""
Builds the QA test-case catalogue from reviewable YAML sources.

Sources (authoritative, text, reviewed in pull requests):
    qa/<module>/functionalities.yaml     functionality registry (traceability rows)
    qa/<module>/cases/*.yaml             manual test cases (one file == one worksheet)
    qa/<module>/clarifications.yaml      clarification / QA-risk register
    qa/tools/existing_test_layers.yaml   classification of existing xUnit / Vitest files by layer

Generated (never hand-edit):
    qa/HRMS_Test_Cases.xlsx
    qa/<module>/*.generated.csv          flat step-level export and traceability
    qa/<module>/coverage-stats.json      numbers used by the coverage review

Validation performed (build fails on any error):
    * required fields present, controlled vocabularies respected, ids unique
    * every `func` id exists in the registry; every functionality has at least one case
    * every referenced existing test (xU:/V:) really exists in the test sources
    * no secret-looking material in any source text

Usage:  python qa/tools/build_test_catalogue.py [--module 01-authentication]
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
QA = ROOT / "qa"
BACKEND_TESTS = ROOT / "Backend" / "HRMS.Tests"
FRONTEND_SRC = ROOT / "Frontend" / "HRMS.Web" / "src"

TYPES = ["Positive", "Negative", "Boundary", "Security", "Authorization", "Tenant-Isolation",
         "Integration-Concurrency", "Contract"]
TAGS = TYPES + ["Validation", "Idempotency", "Configuration", "Usability", "Environment"]
PRIORITIES = ["P0", "P1", "P2", "P3"]
SEVERITIES = ["Critical", "High", "Medium", "Low"]
AUTO = ["API", "UI", "API + UI", "DB-assisted", "Existing automated coverage", "Manual only"]
STATUSES = ["Not Run", "Pass", "Fail", "Blocked", "Skipped",
            "NOT EXECUTED - environment unavailable"]

SECRET_PATTERNS = [
    re.compile(r"eyJ[A-Za-z0-9_-]{15,}\.[A-Za-z0-9_-]{10,}"),              # JWT-looking
    re.compile(r"(?i)\b(password|secret|pwd)\s*[:=]\s*[\"']?[A-Za-z0-9!@#$%^&*]{6,}"),
    re.compile(r"(?i)server=.+;.*(password|pwd)="),                          # connection string
]

STEP_API = re.compile(r"^\[(GET|POST|PUT|DELETE|PATCH|OPTIONS|HEAD)\s+([^\]\s]+)\]\s*")
EXP_STATUS = re.compile(r"^\[(\d{3}(?:/\d{3})*)\]\s*")

COLUMNS = [
    "Test Case ID", "Module", "Sub Module", "Functionality", "Scenario", "Test Case Title",
    "Requirement/Source Reference", "Test Type", "Priority", "Severity if Failed", "Preconditions",
    "User Role", "Required Permission", "Tenant", "Test Data", "Step No", "Test Step",
    "Expected Result", "API Endpoint", "HTTP Method", "Expected HTTP Status", "Database Validation",
    "UI Validation", "Security Validation", "Automation Candidate", "Recommended Automation Layer",
    "Smoke Candidate", "Regression Candidate", "Security Candidate", "Existing Automated Coverage",
    "Existing Coverage Layer", "Clarification Ref", "Actual Result", "Status", "Defect ID", "Comments",
]
FIRST_ROW_ONLY = {"Requirement/Source Reference", "Preconditions", "Test Data", "Database Validation",
                  "UI Validation", "Security Validation", "Existing Automated Coverage",
                  "Existing Coverage Layer", "Clarification Ref", "Comments", "Scenario"}


class BuildError(Exception):
    pass


def load_yaml(path: Path):
    with path.open(encoding="utf-8") as handle:
        return yaml.safe_load(handle)


def as_list(value):
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


# ----------------------------------------------------------------------------- existing test verification
class ExistingTests:
    def __init__(self, layers_path: Path):
        cfg = load_yaml(layers_path)
        self.xunit_layers = cfg["xunit"]
        self.vitest_layers = cfg["vitest"]
        self._xunit_cache: dict[str, str] = {}
        self._vitest_cache: dict[str, str] = {}

    def _xunit_source(self, cls: str) -> str | None:
        if cls not in self._xunit_cache:
            hits = list(BACKEND_TESTS.rglob(f"{cls}.cs"))
            self._xunit_cache[cls] = hits[0].read_text(encoding="utf-8-sig") if hits else ""
        return self._xunit_cache[cls] or None

    def _vitest_source(self, rel: str) -> str | None:
        if rel not in self._vitest_cache:
            path = FRONTEND_SRC / rel
            self._vitest_cache[rel] = path.read_text(encoding="utf-8") if path.exists() else ""
        return self._vitest_cache[rel] or None

    def resolve(self, ref: str) -> tuple[str, str]:
        """Returns (display, layer). Raises BuildError when the referenced test does not exist."""
        if ref.startswith("xU:"):
            body = ref[3:]
            cls, _, method = body.partition(".")
            source = self._xunit_source(cls)
            if source is None:
                raise BuildError(f"existing xUnit class not found: {cls}")
            if method and not re.search(rf"\b{re.escape(method)}\s*\(", source):
                raise BuildError(f"existing xUnit method not found: {body}")
            layer = self.xunit_layers.get(cls)
            if layer is None:
                raise BuildError(f"no layer classification for xUnit class {cls}")
            return f"xUnit {body}", layer
        if ref.startswith("V:"):
            body = ref[2:]
            rel, _, name = body.partition("::")
            source = self._vitest_source(rel)
            if source is None:
                raise BuildError(f"existing Vitest file not found: {rel}")
            if name and name not in source:
                raise BuildError(f"existing Vitest case not found in {rel}: {name}")
            layer = self.vitest_layers.get(rel)
            if layer is None:
                raise BuildError(f"no layer classification for Vitest file {rel}")
            return f"Vitest {rel} :: {name}" if name else f"Vitest {rel}", layer
        raise BuildError(f"unknown existing-test reference prefix: {ref}")


# ----------------------------------------------------------------------------- case model
def normalise_case(raw: dict, defaults: dict, sheet: str, source_file: str) -> dict:
    case = {**defaults, **raw}
    case["_sheet"] = sheet
    case["_file"] = source_file
    return case


def validate_case(case: dict, func_ids: set[str], existing: ExistingTests, errors: list[str]):
    cid = case.get("id", "<missing id>")

    def err(msg):
        errors.append(f"{cid}: {msg}")

    for key in ("id", "func", "sc", "t", "src", "type", "pri", "sev", "pre", "role", "perm", "tenant",
                "data", "steps", "auto"):
        if case.get(key) in (None, "", []):
            err(f"missing required field '{key}'")
    if case.get("type") not in TYPES:
        err(f"type '{case.get('type')}' not in {TYPES}")
    for tag in as_list(case.get("tags")):
        if tag not in TAGS:
            err(f"tag '{tag}' not in {TAGS}")
    if case.get("pri") not in PRIORITIES:
        err(f"priority '{case.get('pri')}' invalid")
    if case.get("sev") not in SEVERITIES:
        err(f"severity '{case.get('sev')}' invalid")
    if case.get("auto") not in AUTO:
        err(f"automation classification '{case.get('auto')}' not in {AUTO}")
    for f in as_list(case.get("func")):
        if f not in func_ids:
            err(f"unknown functionality id {f}")
    steps = as_list(case.get("steps"))
    if not steps:
        err("no steps")
    for idx, step in enumerate(steps, 1):
        if not isinstance(step, str) or " => " not in step:
            err(f"step {idx} must be a string 'step => expected'")
    for ref in as_list(case.get("ex")):
        try:
            existing.resolve(ref)
        except BuildError as exc:
            err(str(exc))
    if case.get("auto") == "Existing automated coverage" and not as_list(case.get("ex")):
        err("classified 'Existing automated coverage' but lists no existing test")
    blob = json.dumps(case, default=str)
    for pattern in SECRET_PATTERNS:
        if pattern.search(blob):
            err(f"possible secret material matched {pattern.pattern[:30]}...")


def split_step(step: str):
    text, _, expected = step.partition(" => ")
    method = path = status = ""
    m = STEP_API.match(text)
    if m:
        method, path = m.group(1), m.group(2)
        text = text[m.end():]
    m2 = EXP_STATUS.match(expected)
    if m2:
        status = m2.group(1)
        expected = expected[m2.end():]
    return text.strip(), expected.strip(), method, path, status


def case_rows(case: dict, func_index: dict, existing: ExistingTests, module: str):
    rows = []
    funcs = as_list(case["func"])
    func_label = "; ".join(f"{f} {func_index[f]['name']}" for f in funcs)
    refs = as_list(case.get("ex"))
    displays, layers = [], []
    for ref in refs:
        display, layer = existing.resolve(ref)
        displays.append(display)
        layers.append(layer)
    first_api = first_method = first_status = ""
    parsed = [split_step(s) for s in as_list(case["steps"])]
    for text, expected, method, path, status in parsed:
        if path and not first_api:
            first_api, first_method = path, method
        if status and not first_status:
            first_status = status
    for idx, (text, expected, method, path, status) in enumerate(parsed, 1):
        row = OrderedDict.fromkeys(COLUMNS, "")
        row.update({
            "Test Case ID": case["id"], "Module": module, "Sub Module": case.get("sub", ""),
            "Functionality": func_label, "Scenario": case["sc"], "Test Case Title": case["t"],
            "Requirement/Source Reference": "; ".join(as_list(case["src"])),
            "Test Type": case["type"] + (" (+ " + ", ".join(t for t in as_list(case.get("tags"))
                                                            if t != case["type"]) + ")"
                                        if [t for t in as_list(case.get("tags")) if t != case["type"]] else ""),
            "Priority": case["pri"], "Severity if Failed": case["sev"],
            "Preconditions": case["pre"], "User Role": case["role"], "Required Permission": case["perm"],
            "Tenant": case["tenant"], "Test Data": case["data"], "Step No": idx, "Test Step": text,
            "Expected Result": expected, "API Endpoint": path or (first_api if idx == 1 else ""),
            "HTTP Method": method or (first_method if idx == 1 and not path else ""),
            "Expected HTTP Status": status or (first_status if idx == 1 and not status and not path else ""),
            "Database Validation": case.get("db", "N/A"), "UI Validation": case.get("ui", "N/A"),
            "Security Validation": case.get("sec", "N/A"),
            "Automation Candidate": "No" if case["auto"] == "Manual only" else "Yes",
            "Recommended Automation Layer": case["auto"],
            "Smoke Candidate": "Yes" if case.get("smoke") else "No",
            "Regression Candidate": "Yes" if case.get("reg", True) else "No",
            "Security Candidate": "Yes" if case.get("secc") or case["type"] in ("Security", "Authorization",
                                                                             "Tenant-Isolation") else "No",
            "Existing Automated Coverage": "\n".join(displays) if displays else "None",
            "Existing Coverage Layer": "\n".join(sorted(set(layers))) if layers else "None",
            "Clarification Ref": ", ".join(as_list(case.get("cr"))),
            "Actual Result": "", "Status": "Not Run", "Defect ID": "", "Comments": case.get("cm", ""),
        })
        rows.append(row)
    return rows


# ----------------------------------------------------------------------------- workbook helpers
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WIDTHS = {"Test Case ID": 18, "Module": 18, "Sub Module": 20, "Functionality": 34, "Scenario": 34,
          "Test Case Title": 44, "Requirement/Source Reference": 34, "Test Type": 18, "Priority": 8,
          "Severity if Failed": 12, "Preconditions": 40, "User Role": 18, "Required Permission": 22,
          "Tenant": 10, "Test Data": 34, "Step No": 6, "Test Step": 56, "Expected Result": 56,
          "API Endpoint": 30, "HTTP Method": 9, "Expected HTTP Status": 10, "Database Validation": 34,
          "UI Validation": 30, "Security Validation": 34, "Automation Candidate": 11,
          "Recommended Automation Layer": 20, "Smoke Candidate": 8, "Regression Candidate": 10,
          "Security Candidate": 9, "Existing Automated Coverage": 46, "Existing Coverage Layer": 28,
          "Clarification Ref": 12, "Actual Result": 28, "Status": 16, "Defect ID": 12, "Comments": 36}


def write_table(ws, headers, rows, widths=None, wrap=True):
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") if isinstance(row, dict) else row[i] for i, h in enumerate(headers)]
                  if isinstance(row, dict) else list(row))
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = (widths or {}).get(header, 22)
    if wrap:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def add_dropdown(ws, headers, header_name, choices, allow_blank=False):
    if header_name not in headers or ws.max_row < 2:
        return
    col = get_column_letter(headers.index(header_name) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=allow_blank)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{ws.max_row + 500}")


def add_status_validation(ws, headers):
    add_dropdown(ws, headers, "Status", STATUSES)
    add_dropdown(ws, headers, "Priority", PRIORITIES)
    add_dropdown(ws, headers, "Severity if Failed", SEVERITIES)


# ----------------------------------------------------------------------------- main build
def build_module(module_dir: Path, existing: ExistingTests):
    module_name = "Authentication & Tenant Security" if module_dir.name == "01-authentication" else module_dir.name
    func_doc = load_yaml(module_dir / "functionalities.yaml")
    functionalities = func_doc["functionalities"]
    func_index = {f["id"]: f for f in functionalities}
    if len(func_index) != len(functionalities):
        raise BuildError("duplicate functionality ids")

    errors: list[str] = []
    sheets: "OrderedDict[str, list[dict]]" = OrderedDict()
    all_cases: list[dict] = []
    for path in sorted((module_dir / "cases").glob("*.yaml")):
        doc = load_yaml(path)
        sheet = doc["sheet"]
        defaults = {"sub": doc.get("submodule", ""), **doc.get("defaults", {})}
        for raw in doc["cases"]:
            case = normalise_case(raw, defaults, sheet, path.name)
            validate_case(case, set(func_index), existing, errors)
            sheets.setdefault(sheet, []).append(case)
            all_cases.append(case)

    ids = Counter(c.get("id") for c in all_cases)
    for cid, n in ids.items():
        if n > 1:
            errors.append(f"{cid}: duplicate test case id")
    covered = {f for c in all_cases for f in as_list(c.get("func"))}
    for fid in func_index:
        if fid not in covered:
            errors.append(f"{fid}: functionality has no test case")
    if errors:
        raise BuildError("\n".join(errors))

    clar_path = module_dir / "clarifications.yaml"
    clarifications = load_yaml(clar_path)["clarifications"] if clar_path.exists() else []
    known_cr = {c["id"] for c in clarifications}
    for c in all_cases:
        for cr in as_list(c.get("cr")):
            if cr not in known_cr:
                errors.append(f"{c['id']}: unknown clarification {cr}")
    if errors:
        raise BuildError("\n".join(errors))

    return module_name, func_index, sheets, all_cases, clarifications


def dims(case) -> set[str]:
    return {case["type"], *as_list(case.get("tags"))}


def stats_for(all_cases, func_index):
    by_type = Counter(c["type"] for c in all_cases)
    by_auto = Counter(c["auto"] for c in all_cases)
    by_pri = Counter(c["pri"] for c in all_cases)
    by_sev = Counter(c["sev"] for c in all_cases)
    tag_counter = Counter(t for c in all_cases for t in dims(c))
    layer_counter = Counter()
    for c in all_cases:
        a = c["auto"]
        if a in ("API", "API + UI", "DB-assisted"):
            layer_counter["API-capable (API / API + UI / DB-assisted)"] += 1
        if a in ("UI", "API + UI"):
            layer_counter["UI-capable (UI / API + UI)"] += 1
    flags = {
        "smoke": sum(1 for c in all_cases if c.get("smoke")),
        "regression": sum(1 for c in all_cases if c.get("reg", True)),
        "security_candidate": sum(1 for c in all_cases
                                  if c.get("secc") or c["type"] in ("Security", "Authorization", "Tenant-Isolation")),
        "with_existing_mapping": sum(1 for c in all_cases if as_list(c.get("ex"))),
        "without_existing_mapping": sum(1 for c in all_cases if not as_list(c.get("ex"))),
    }
    dim_cols = ["Positive", "Negative", "Validation", "Boundary", "Authorization", "Tenant-Isolation",
                "Security", "Integration-Concurrency", "Contract"]
    matrix = {}
    for fid in func_index:
        related = [c for c in all_cases if fid in as_list(c["func"])]
        row = {"cases": len(related)}
        for d in dim_cols:
            row[d] = sum(1 for c in related if d in dims(c))
        row["API"] = sum(1 for c in related if c["auto"] in ("API", "API + UI", "DB-assisted"))
        row["UI"] = sum(1 for c in related if c["auto"] in ("UI", "API + UI"))
        row["ExistingMapped"] = sum(1 for c in related if as_list(c.get("ex")))
        matrix[fid] = row
    return {
        "total_cases": len(all_cases), "total_functionalities": len(func_index),
        "by_primary_type": dict(by_type), "by_automation_classification": dict(by_auto),
        "by_priority": dict(by_pri), "by_severity": dict(by_sev), "dimension_counts_incl_tags": dict(tag_counter),
        "layer_capability": dict(layer_counter), "flags": flags, "coverage_matrix": matrix,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--module", action="append", help="module directory name under qa/")
    args = parser.parse_args()
    modules = args.module or sorted(p.name for p in QA.iterdir()
                                    if p.is_dir() and re.match(r"\d\d-", p.name) and (p / "cases").exists())
    existing = ExistingTests(QA / "tools" / "existing_test_layers.yaml")
    wb = Workbook()
    wb.remove(wb.active)
    summary_rows, trace_ws_rows, exist_rows, clar_rows = [], [], [], []

    for module in modules:
        module_dir = QA / module
        try:
            module_name, func_index, sheets, all_cases, clarifications = build_module(module_dir, existing)
        except BuildError as exc:
            print("BUILD FAILED for", module, "\n" + str(exc), file=sys.stderr)
            return 1
        stats = stats_for(all_cases, func_index)
        (module_dir / "coverage-stats.json").write_text(json.dumps(stats, indent=2), encoding="utf-8")

        flat_rows = []
        for sheet_name, cases in sheets.items():
            ws = wb.create_sheet(sheet_name[:31])
            rows = []
            for case in cases:
                rows.extend(case_rows(case, func_index, existing, module_name))
            flat_rows.extend(rows)
            write_table(ws, COLUMNS, rows, WIDTHS)
            add_status_validation(ws, COLUMNS)
        # flat CSV for reviewers who do not open Excel. Named after the module directory (numeric
        # prefix stripped) so each module gets its own file instead of every module colliding on
        # a name hardcoded for Module 1.
        slug = re.sub(r"^\d+-", "", module)
        with (module_dir / f"{slug}-test-cases.generated.csv").open("w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows(flat_rows)

        # traceability
        for fid, f in func_index.items():
            related = [c["id"] for c in all_cases if fid in as_list(c["func"])]
            ex_refs = sorted({existing.resolve(r)[0] for c in all_cases if fid in as_list(c["func"])
                              for r in as_list(c.get("ex"))})
            trace_ws_rows.append([module_name, f["group"], fid, f["name"], f["status"], f.get("ui", ""),
                                  f.get("api", ""), f.get("service", ""), f.get("permission", ""),
                                  f.get("db", ""), len(related), ", ".join(related),
                                  len(ex_refs), "\n".join(ex_refs), "", ""])
        for c in all_cases:
            for ref in as_list(c.get("ex")):
                display, layer = existing.resolve(ref)
                exist_rows.append([c["id"], c["t"], display, layer,
                                   "Yes" if c["auto"] == "Existing automated coverage" else "Partial / supporting"])
        for cl in clarifications:
            related = [c["id"] for c in all_cases if cl["id"] in as_list(c.get("cr"))]
            clar_rows.append([cl["id"], cl["title"], cl.get("observation", ""), cl.get("risk", ""),
                              cl.get("source", ""), cl.get("classification", "Open - awaiting decision"),
                              ", ".join(related)])
        with (module_dir / f"{slug}-traceability.generated.csv").open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            writer.writerow(["Module", "Group", "Functionality ID", "Functionality", "Implementation Status",
                             "UI", "API", "Service", "Permission", "DB Impact", "Manual Case Count",
                             "Manual Test Cases", "Existing Test Count", "Existing Tests",
                             "Automation Tests (Playwright)", "Defects"])
            writer.writerows(r for r in trace_ws_rows if r[0] == module_name)

        for t, n in sorted(stats["by_primary_type"].items()):
            summary_rows.append([module_name, "Primary type", t, n])
        for t, n in sorted(stats["by_automation_classification"].items()):
            summary_rows.append([module_name, "Automation classification", t, n])
        for t, n in sorted(stats["by_priority"].items()):
            summary_rows.append([module_name, "Priority", t, n])
        summary_rows.append([module_name, "Totals", "Test cases", stats["total_cases"]])
        summary_rows.append([module_name, "Totals", "Functionalities", stats["total_functionalities"]])
        print(f"{module}: {stats['total_cases']} cases, {stats['total_functionalities']} functionalities, "
              f"{sum(len(c['steps']) for c in all_cases)} steps")

    ws = wb.create_sheet("Test_Case_Summary", 0)
    write_table(ws, ["Module", "Dimension", "Value", "Count"], summary_rows,
                {"Module": 34, "Dimension": 26, "Value": 40, "Count": 10})
    ws = wb.create_sheet("Traceability")
    trace_headers = ["Module", "Group", "Functionality ID", "Functionality", "Implementation Status", "UI", "API",
                     "Service", "Permission", "DB Impact", "Manual Case Count", "Manual Test Cases",
                     "Existing Test Count", "Existing Tests", "Automation Tests (Playwright)", "Defects"]
    write_table(ws, trace_headers, trace_ws_rows,
                {"Functionality": 40, "UI": 30, "API": 40, "Service": 36, "Permission": 26, "DB Impact": 36,
                 "Manual Test Cases": 50, "Existing Tests": 60})
    ws = wb.create_sheet("Existing_Test_Map")
    write_table(ws, ["Manual Test Case", "Title", "Existing Test", "Existing Layer (NOT live-stack)",
                     "Equivalence"], exist_rows, {"Title": 50, "Existing Test": 70,
                                                  "Existing Layer (NOT live-stack)": 40})
    ws = wb.create_sheet("Clarifications")
    write_table(ws, ["ID", "Title", "Observation (current behaviour)", "Risk", "Source",
                     "Classification (Expected / Confirmed defect / Design change / Accepted risk)",
                     "Exposing Test Cases"], clar_rows,
                {"Title": 40, "Observation (current behaviour)": 70, "Risk": 50, "Source": 40,
                 "Classification (Expected / Confirmed defect / Design change / Accepted risk)": 30,
                 "Exposing Test Cases": 40})
    ws = wb.create_sheet("Defects")
    write_table(ws, ["Defect ID", "Title", "Module", "Functionality", "Test Case ID", "Environment",
                     "Build/Commit", "Severity", "Priority", "Preconditions", "Steps to Reproduce",
                     "Expected Result", "Actual Result", "Evidence", "API Request", "API Response",
                     "Screenshot Path", "Trace Path", "Database Evidence", "Status", "Assigned To",
                     "Found Date", "Retest Date", "Retest Result", "Comments"], [])
    ws = wb.create_sheet("Execution_Summary")
    write_table(ws, ["Execution ID", "Environment", "Date/Time", "Git Commit SHA", "Browser",
                     "Database Provider", "Total", "Executed", "Passed", "Failed", "Blocked", "Skipped",
                     "NOT EXECUTED - environment unavailable", "Pass %", "New Defects", "Reopened Defects",
                     "Known Defects", "Critical Failures"], [])
    wb.save(QA / "HRMS_Test_Cases.xlsx")
    print("wrote", (QA / "HRMS_Test_Cases.xlsx").relative_to(ROOT))
    return 0


if __name__ == "__main__":
    sys.exit(main())
